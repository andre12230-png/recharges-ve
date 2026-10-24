"""Le texte blanc des en-tetes de tableau du PDF se lit (relecture du
14/10/2026).

L'en-tete du premier tableau (synthese) etait d'un bleu clair, #3A7BD5 :
le blanc n'y faisait que 4,22 pour 1, sous les 4,5 de la charte. Les autres
en-tetes, en #2A6BC5, font 5,24. Ce test lit chaque fond d'en-tete ecrit en
blanc dans export.py.
"""
from __future__ import annotations

import re
from pathlib import Path

RACINE = Path(__file__).resolve().parent.parent


def _luminance(couleur: str) -> float:
    def canal(v: int) -> float:
        x = v / 255
        return x / 12.92 if x <= 0.03928 else ((x + 0.055) / 1.055) ** 2.4
    r, g, b = (int(couleur[i:i + 2], 16) for i in (1, 3, 5))
    return 0.2126 * canal(r) + 0.7152 * canal(g) + 0.0722 * canal(b)


def _contraste(a: str, b: str) -> float:
    la, lb = _luminance(a), _luminance(b)
    return (max(la, lb) + 0.05) / (min(la, lb) + 0.05)


def test_texte_blanc_des_en_tetes_lisible():
    source = (RACINE / "export.py").read_text(encoding="utf-8")
    fonds = re.findall(
        r'\("BACKGROUND", \(0, 0\), \(-1, 0\), colors\.HexColor\("(#[0-9A-Fa-f]{6})"\)\),'
        r'\s*\("TEXTCOLOR", \(0, 0\), \(-1, 0\), colors\.white\)',
        source)
    assert len(fonds) == 3, fonds          # synthese, tarifs, detail
    for fond in fonds:
        assert _contraste("#FFFFFF", fond) >= 4.5, fond


def test_couleurs_de_l_excel_lisibles(tmp_path):
    """L'Excel utilisait le meme bleu clair : en-tete blanc sur #3A7BD5 et
    titre de la Synthese en #3A7BD5 sur le blanc, 4,22 pour 1."""
    from openpyxl import load_workbook

    import export
    from test_export_complet import RECHARGES
    chemin = tmp_path / "essai.xlsx"
    export.exporter_excel(RECHARGES, chemin, "Tout")
    classeur = load_workbook(chemin)
    entete = classeur.worksheets[0].cell(row=1, column=1)
    fond = "#" + entete.fill.fgColor.rgb[-6:]
    assert _contraste("#FFFFFF", fond) >= 4.5, fond
    titre = classeur["Synthèse"]["A1"]
    encre = "#" + titre.font.color.rgb[-6:]
    assert _contraste(encre, "#FFFFFF") >= 4.5, encre


def test_titres_du_pdf_lisibles_sur_le_blanc():
    """Le titre et les sous-titres s'ecrivent en couleur sur le blanc de la
    page : le titre etait lui aussi en #3A7BD5 (4,22 pour 1)."""
    source = (RACINE / "export.py").read_text(encoding="utf-8")
    encres = re.findall(
        r'ParagraphStyle\("\w+", parent=styles\["\w+"\],\s*'
        r'textColor=colors\.HexColor\("(#[0-9A-Fa-f]{6})"\)\)', source)
    assert len(encres) == 2, encres        # titre, sous-titres
    for encre in encres:
        assert _contraste(encre, "#FFFFFF") >= 4.5, encre
