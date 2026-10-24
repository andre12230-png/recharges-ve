"""Frais des cartes de recharge (28/09/2026).

Achat d'une carte (une fois) ou abonnement mensuel a un reseau : comptes a
part, comme l'abonnement de la maison, jamais dans le cout des recharges ni
dans le prix du kWh. Enregistres dans tarifs.json, cle « frais_cartes ».
Donnees inventees, dans le dossier temporaire de tests/conftest.py.

    py -m pytest tests/test_frais_cartes.py
"""

from __future__ import annotations

import json
from datetime import date

import pytest

import recharges as r
from saisie import ErreurSaisie, lire_frais_carte

ACHAT = {"carte": "Chargemap", "type": "unique", "date": "2025-03-10",
         "fin": None, "montant": 20.0}
ABONNEMENT = {"carte": "Réseau B", "type": "mensuel", "date": "2025-04-01",
              "fin": "2025-06-30", "montant": 6.0}


@pytest.fixture
def fichier():
    """Un tarifs.json neuf, avec une grille et un reglage a conserver."""
    r.TARIFS_FILE.write_text(json.dumps(
        {"tarifs": [], "capacite_batterie": 50}), encoding="utf-8")
    yield r.TARIFS_FILE


# ---------- le calcul (sans fichier) ----------

def test_un_achat_compte_le_jour_ou_il_est_fait():
    f = [ACHAT]
    assert r.calculer_frais_cartes(date(2025, 3, 1), date(2025, 3, 31),
                                   f)["total"] == 20.0
    assert r.calculer_frais_cartes(date(2025, 4, 1), date(2025, 12, 31),
                                   f)["total"] == 0.0


def test_un_abonnement_compte_chaque_mois_au_prorata():
    f = [ABONNEMENT]
    # Avril, mai, juin entiers : 3 x 6 EUR.
    assert r.calculer_frais_cartes(date(2025, 1, 1), date(2025, 12, 31),
                                   f)["total"] == 18.0
    # La moitie de juin (15 jours sur 30).
    assert r.calculer_frais_cartes(date(2025, 6, 16), date(2025, 6, 30),
                                   f)["total"] == 3.0


def test_un_abonnement_sans_fin_court_jusqu_a_la_fin_demandee():
    f = [dict(ABONNEMENT, fin=None)]
    assert r.calculer_frais_cartes(date(2025, 4, 1), date(2025, 7, 31),
                                   f)["total"] == 24.0


def test_total_par_carte():
    res = r.calculer_frais_cartes(date(2025, 1, 1), date(2025, 12, 31),
                                  [ACHAT, ABONNEMENT])
    assert res == {"total": 38.0,
                   "par_carte": {"Chargemap": 20.0, "Réseau B": 18.0}}


# ---------- l'enregistrement dans tarifs.json ----------

def test_ajouter_et_supprimer_sans_toucher_au_reste(fichier):
    ajoute = r.ajouter_frais_carte(ACHAT)
    assert ajoute["id"]
    contenu = json.loads(fichier.read_text(encoding="utf-8"))
    assert contenu["capacite_batterie"] == 50 and contenu["tarifs"] == []
    assert [f["carte"] for f in r.charger_frais_cartes()] == ["Chargemap"]
    assert r.supprimer_frais_carte(ajoute["id"]) is True
    assert r.charger_frais_cartes() == []
    assert r.supprimer_frais_carte("inconnu") is False


def test_un_ancien_fichier_sans_frais(fichier):
    """Les fichiers d'avant le 28/09/2026 n'ont pas la cle : aucun frais."""
    assert r.charger_frais_cartes() == []


def test_la_grille_enregistree_garde_les_frais(fichier):
    """Enregistrer un tarif reecrit tarifs.json : les frais doivent rester."""
    r.ajouter_frais_carte(ACHAT)
    r.enregistrer_tarif({"nom": "Contrat", "debut": "2025-01-01", "fin": None,
                         "type": "simple", "prix_kwh": 0.2,
                         "abonnement_mensuel": 10.0})
    assert len(r.charger_frais_cartes()) == 1


def test_le_moteur_refuse_des_frais_impossibles(fichier):
    with pytest.raises(ValueError):
        r.ajouter_frais_carte(dict(ACHAT, montant=0))
    with pytest.raises(ValueError):
        r.ajouter_frais_carte(dict(ABONNEMENT, fin="2025-01-01"))
    with pytest.raises(ValueError):
        r.ajouter_frais_carte(dict(ACHAT, type="hebdomadaire"))
    assert r.charger_frais_cartes() == []


# ---------- la saisie ----------

def textes(**modifs):
    base = {"carte": "Chargemap", "type": "unique", "date": "10/03/2025",
            "fin": "", "montant": "20,00"}
    base.update(modifs)
    return base


def test_lire_un_achat():
    assert lire_frais_carte(textes()) == ACHAT


def test_lire_un_abonnement():
    lu = lire_frais_carte(textes(carte=" Réseau B ", type="mensuel",
                                 date="01/04/2025", fin="30/06/2025",
                                 montant="6"))
    assert lu == ABONNEMENT


def test_la_fin_d_un_achat_est_ignoree():
    assert lire_frais_carte(textes(fin="30/06/2025"))["fin"] is None


@pytest.mark.parametrize("modifs, champ", [
    ({"carte": "  "}, "carte"),
    ({"montant": ""}, "montant"),
    ({"montant": "0"}, "montant"),
    ({"date": "31/02/2025"}, "date"),
    ({"type": "mensuel", "fin": "01/01/2025"}, "fin"),
])
def test_les_saisies_refusees(modifs, champ):
    with pytest.raises(ErreurSaisie) as e:
        lire_frais_carte(textes(**modifs))
    assert e.value.champ == champ


# ---------- la page Tarifs ----------

@pytest.fixture
def fenetre(qapp, monkeypatch, fichier):
    import fenetre as module_fenetre
    monkeypatch.setattr(module_fenetre, "charger_donnees",
                        lambda: {"recharges": []})
    f = module_fenetre.FenetrePrincipale()
    yield f
    f.deleteLater()


def test_la_page_tarifs_ajoute_des_frais(fenetre, qapp):
    from pages import tarifs as ta
    page = ta.construire(fenetre)
    assert page.lignes_frais == []
    cases = page.cases_frais
    cases["carte"].setCurrentText("Chargemap")
    cases["date"].setText("10/03/2025")
    cases["montant"].setText("20")
    page.bouton_frais.click()
    assert [f["montant"] for f in r.charger_frais_cartes()] == [20.0]
    # La page se reconstruit avec la nouvelle ligne et le message de succes,
    # garde sur la fenetre le temps de la reconstruction.
    assert fenetre.message_frais.startswith("Frais enregistrés")
    page = ta.construire(fenetre)
    assert len(page.lignes_frais) == 1
    assert page.message_frais.text().startswith("Frais enregistrés")


def test_la_page_tarifs_refuse_une_saisie(fenetre):
    from pages import tarifs as ta
    from theme_qt import ROUGE
    page = ta.construire(fenetre)
    page.cases_frais["carte"].setCurrentText("Chargemap")
    page.cases_frais["date"].setText("10/03/2025")
    page.cases_frais["montant"].setText("douze")
    page.bouton_frais.click()
    assert ROUGE in page.message_frais.styleSheet()
    assert "n'est pas un nombre" in page.message_frais.text()
    assert r.charger_frais_cartes() == []


def test_la_fin_ne_sert_qu_a_un_abonnement(fenetre):
    from pages import tarifs as ta
    page = ta.construire(fenetre)
    assert not page.cases_frais["fin"].isEnabled()
    page.cases_frais["type"].setCurrentIndex(1)       # Abonnement mensuel
    assert page.cases_frais["fin"].isEnabled()


def test_supprimer_des_frais(fenetre, monkeypatch):

    from pages import tarifs as ta
    r.ajouter_frais_carte(ACHAT)
    monkeypatch.setattr(ta, "confirmer", lambda *a, **k: True)
    page = ta.construire(fenetre)
    page.lignes_frais[0].bouton.click()
    assert r.charger_frais_cartes() == []


# ---------- le tableau de bord et les exports ----------

def test_le_bloc_exterieur_montre_les_frais(fichier):
    from pages import tableau_de_bord as tb
    r.ajouter_frais_carte(ACHAT)
    c = tb.chiffres_de_la_periode([], "Mars 2025")
    assert c["exterieur"]["frais_cartes"] == 20.0
    assert tb.chiffres_de_la_periode([], "Avril 2025")[
        "exterieur"]["frais_cartes"] == 0.0


def test_les_exports_donnent_les_frais(fichier, tmp_path):
    from openpyxl import load_workbook

    import export
    from pages import exporter as ex
    r.ajouter_frais_carte(ACHAT)
    recharge = {"id": "x", "date": "2025-03-12T10:00", "lieu": "exterieur",
                "kwh": 10.0, "cout_eur": 5.0, "operateur": "Ionity"}
    c = ex.complements([recharge], "Extérieur", "2025", conso={
        "intervalles": []})
    assert c["frais_cartes"] == 20.0
    assert ex.complements([recharge], "Maison", "2025",
                          conso={"intervalles": []})["frais_cartes"] is None
    chemin = tmp_path / "x.xlsx"
    export.exporter_excel([recharge], chemin, "2025", complements=c)
    lignes = {rang[0]: rang[1] for rang in load_workbook(chemin)["Synthèse"]
              .iter_rows(values_only=True) if rang[0]}
    assert lignes["Frais de cartes de recharge de la période (EUR)"] == 20.0


def test_tout_compte_une_carte_achetee_avant_la_premiere_recharge(fichier):
    """Defaut trouve le 28/09/2026 sur les vraies donnees : « Tout » commence
    a la premiere recharge, et une carte achetee AVANT (on l'achete pour
    s'en servir) n'etait pas comptee."""
    from pages import exporter as ex
    from pages import tableau_de_bord as tb
    r.ajouter_frais_carte(ACHAT)                    # 10/03/2025
    recharge = {"id": "x", "date": "2025-07-01T10:00", "lieu": "exterieur",
                "kwh": 10.0, "cout_eur": 5.0, "operateur": "Ionity"}
    assert tb.chiffres_de_la_periode([recharge], "Tout")[
        "exterieur"]["frais_cartes"] == 20.0
    assert ex.complements([recharge], "Tout", "Tout", conso={
        "intervalles": []})["frais_cartes"] == 20.0
    # Sur une annee, elle reste a sa date : rien en 2026.
    assert tb.chiffres_de_la_periode([recharge], "2026")[
        "exterieur"]["frais_cartes"] == 0.0


# ---------- abonnement annuel (28/09/2026) ----------
# Le prix de l'annee compte en entier le jour du debut, puis a chaque date
# anniversaire tant que l'abonnement court : ce qui sort du compte.

ANNUEL = {"carte": "Réseau C", "type": "annuel", "date": "2024-02-15",
          "fin": None, "montant": 60.0}


def test_annuel_compte_le_debut_puis_chaque_anniversaire():
    f = [ANNUEL]
    total = lambda d, fin: r.calculer_frais_cartes(d, fin, f)["total"]  # noqa: E731
    assert total(date(2024, 1, 1), date(2024, 12, 31)) == 60.0
    assert total(date(2025, 2, 1), date(2025, 2, 28)) == 60.0   # renouvele
    assert total(date(2025, 3, 1), date(2026, 1, 31)) == 0.0
    assert total(date(2024, 1, 1), date(2026, 12, 31)) == 180.0


def test_annuel_s_arrete_a_sa_fin():
    f = [dict(ANNUEL, fin="2025-12-31")]
    # Debut 2024, renouvele en 2025, arrete avant fevrier 2026.
    assert r.calculer_frais_cartes(date(2024, 1, 1), date(2026, 12, 31),
                                   f)["total"] == 120.0


def test_annuel_commence_un_29_fevrier():
    f = [dict(ANNUEL, date="2024-02-29")]
    # En 2025, pas de 29 fevrier : l'anniversaire tombe le 28.
    assert r.calculer_frais_cartes(date(2025, 2, 28), date(2025, 2, 28),
                                   f)["total"] == 60.0


def test_lire_un_abonnement_annuel():
    lu = lire_frais_carte(textes(carte="Réseau C", type="annuel",
                                 date="15/02/2024", fin="", montant="60"))
    assert lu == ANNUEL
    with pytest.raises(ErreurSaisie) as e:
        lire_frais_carte(textes(type="annuel", date="15/02/2024",
                                fin="01/01/2024"))
    assert e.value.champ == "fin"


def test_ligne_d_un_abonnement_annuel():
    from pages import tarifs as ta
    assert ta.ligne_frais(ANNUEL) == (
        "depuis le 15/02/2024  ·  Réseau C  ·  abonnement annuel  ·  "
        "60,00\xa0€ par an")
    assert "abonnement mensuel" in ta.ligne_frais(ABONNEMENT)


def test_le_formulaire_propose_l_annuel(fenetre):
    from pages import tarifs as ta
    page = ta.construire(fenetre)
    choix = page.cases_frais["type"]
    assert [choix.itemText(i) for i in range(choix.count())] == [
        "Achat unique", "Abonnement mensuel", "Abonnement annuel"]
    choix.setCurrentIndex(2)
    assert page.cases_frais["fin"].isEnabled()
