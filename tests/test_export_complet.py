"""Exports Excel et PDF complets (28/09/2026).

Avant, ni l'un ni l'autre ne donnait l'abonnement, les kilometres ou le
cout aux 100 km, et le PDF ne listait pas les recharges : quelqu'un qui
imprimait son annee n'avait que des totaux. Donnees inventees ; les fichiers
sont ecrits dans un dossier jetable.

    py -m pytest tests/test_export_complet.py
"""

from __future__ import annotations

from datetime import date

import pytest
from openpyxl import load_workbook

import export
from pages import exporter as ex

RECHARGES = [
    {"id": "a", "date": "2025-03-10T23:00", "lieu": "maison", "kwh": 20.0,
     "cout_eur": 3.0, "prix_kwh": 0.15, "tarif_nom": "Mon contrat",
     "kwh_hc": 20.0, "kwh_hp": 0.0},
    {"id": "b", "date": "2025-04-02T12:00", "lieu": "exterieur",
     "kwh": 10.0, "cout_eur": 5.0, "prix_kwh": 0.5, "operateur": "Ionity",
     "carte": "", "lieu_detail": "Aire des Tilleuls"},
    {"id": "c", "date": "2026-01-15T22:00", "lieu": "maison", "kwh": 30.0,
     "cout_eur": 4.5, "prix_kwh": 0.15, "tarif_nom": "Mon contrat",
     "kwh_hc": 30.0, "kwh_hp": 0.0},
]

# Trois intervalles entre releves : deux en 2025, un a cheval sur 2025-2026.
CONSO = {"intervalles": [
    {"debut": "2025-03-01T08:00", "fin": "2025-04-01T08:00", "km": 100.0,
     "kwh": 20.0, "cout": 3.0, "n_recharges": 1},
    {"debut": "2025-04-01T08:00", "fin": "2025-05-01T08:00", "km": 300.0,
     "kwh": 10.0, "cout": 5.0, "n_recharges": 1},
    {"debut": "2025-12-01T08:00", "fin": "2026-02-01T08:00", "km": 200.0,
     "kwh": 30.0, "cout": 4.5, "n_recharges": 1},
]}


def _abonnement(debut: date, fin: date) -> dict:
    """Doublure : 10 EUR par mois entame, pour des nombres ronds."""
    mois = (fin.year - debut.year) * 12 + fin.month - debut.month + 1
    return {"total": 10.0 * mois}


def _complements(lieu, periode, conso=CONSO):
    return ex.complements(RECHARGES, lieu, periode, conso=conso,
                          calcul_abonnement=_abonnement)


# ---------- ce qui est ajoute, selon le lieu et la periode ----------

def test_abonnement_de_la_periode():
    assert _complements("Tout", "2025")["abonnement"] == 120.0
    assert _complements("Maison", "Mars 2025")["abonnement"] == 10.0


def test_pas_d_abonnement_pour_l_exterieur_seul():
    """L'abonnement est celui du compteur de la maison : il n'a rien a faire
    dans un export des seules bornes exterieures."""
    assert _complements("Extérieur", "2025")["abonnement"] is None


def test_km_sur_les_seuls_intervalles_compris_dans_la_periode():
    """Comme partout dans l'application, les km se comptent entre deux
    releves : on ne coupe pas un intervalle au 31 decembre, on ne prend que
    ceux qui tiennent entiers dans la periode."""
    km = _complements("Tout", "2025")["km"]
    assert km["km"] == 400.0 and km["cout"] == 8.0
    assert km["cout_100"] == 2.0
    assert km["debut"] == "2025-03-01T08:00" and km["fin"] == "2025-05-01T08:00"


def test_km_sur_tout():
    km = _complements("Tout", "Tout")["km"]
    assert km["km"] == 600.0 and km["cout"] == 12.5


def test_pas_de_km_pour_un_seul_lieu():
    """Le cout aux 100 km vient de TOUTES les recharges de l'intervalle :
    le donner dans un export « Maison » laisserait croire que seules les
    recharges a la maison ont servi."""
    assert _complements("Maison", "2025")["km"] is None


def test_pas_de_km_sans_intervalle_complet_ni_sans_energie():
    assert _complements("Tout", "Mars 2025")["km"] is None
    vide = {"intervalles": [dict(CONSO["intervalles"][0], kwh=0.0, cout=0.0)]}
    # « 0,00 EUR/100 km » ferait croire que rouler ne coute rien.
    assert _complements("Tout", "2025", conso=vide)["km"] is None


# ---------- Excel ----------

def _synthese(tmp_path, complements):
    chemin = tmp_path / "essai.xlsx"
    export.exporter_excel(RECHARGES, chemin, "Tout - 2025",
                          complements=complements)
    feuille = load_workbook(chemin)["Synthèse"]
    return {ligne[0]: ligne[1] for ligne in feuille.iter_rows(values_only=True)
            if ligne[0]}


def test_excel_donne_abonnement_et_km(tmp_path):
    lignes = _synthese(tmp_path, _complements("Tout", "2025"))
    assert lignes["Abonnement électricité de la période (EUR)"] == 120.0
    assert lignes["Distance parcourue (km)"] == 400.0
    assert lignes["Coût aux 100\xa0km (EUR)"] == 2.0


def test_excel_sans_complements_comme_avant(tmp_path):
    lignes = _synthese(tmp_path, None)
    assert "Abonnement électricité de la période (EUR)" not in lignes
    assert "Coût aux 100\xa0km (EUR)" not in lignes
    assert lignes["Nombre de recharges"] == 3


# ---------- PDF ----------

def test_detail_des_recharges_pour_le_pdf():
    lignes = export.lignes_detail(RECHARGES)
    assert lignes[0] == ["Date", "Lieu", "kWh", "Coût (EUR)", "EUR/kWh",
                         "Tarif / Opérateur"]
    assert lignes[1] == ["10/03/2025 23:00", "Maison", "20,00", "3,00",
                         "0,1500", "Mon contrat"]
    assert lignes[2][1] == "Extérieur" and lignes[2][5] == "Ionity"
    assert len(lignes) == 4


def test_pdf_contient_la_liste_et_les_km(tmp_path):
    pypdf = pytest.importorskip("pypdf")
    chemin = tmp_path / "essai.pdf"
    export.exporter_pdf(RECHARGES, chemin, "Tout - 2025",
                        complements=_complements("Tout", "2025"))
    texte = " ".join(p.extract_text() for p in pypdf.PdfReader(chemin).pages)
    assert "Détail des recharges" in texte
    assert "10/03/2025 23:00" in texte and "Aire des Tilleuls" not in texte
    assert "Abonnement" in texte and "120,00" in texte
    assert "2,00 EUR/100 km" in texte


def test_pdf_sans_complements_se_construit(tmp_path):
    chemin = tmp_path / "essai.pdf"
    export.exporter_pdf(RECHARGES, chemin, "Tout")
    assert chemin.stat().st_size > 0


# ---------- la comparaison essence dans les exports (28/09/2026) ----------

ESSENCE = {"km": 1000.0, "litres": 50.0, "cout_essence": 100.0,
           "cout_electrique": 30.0, "economie": 70.0, "cout_100_essence": 10.0,
           "cout_100_electrique": 3.0, "conso_essence": 5.0,
           "prix_essence": 2.0, "debut": "2025-01-01T00:00",
           "fin": "2025-12-31T23:59", "part_non_enregistree": 40,
           "par_annee": []}


def _avec_essence(lieu, periode):
    return ex.complements(RECHARGES, lieu, periode, conso=CONSO,
                          calcul_abonnement=_abonnement,
                          calcul_essence=lambda p: ESSENCE)


def test_la_comparaison_essence_seulement_pour_tous_les_lieux():
    assert _avec_essence("Tout", "2025")["essence"] == ESSENCE
    assert _avec_essence("Maison", "2025")["essence"] is None
    assert _avec_essence("Extérieur", "2025")["essence"] is None


def test_excel_donne_la_comparaison_essence(tmp_path):
    lignes = _synthese(tmp_path, _avec_essence("Tout", "2025"))
    assert lignes["En essence (EUR)"] == 100.0
    assert lignes["Économie (EUR)"] == 70.0


def test_pdf_donne_la_comparaison_essence(tmp_path):
    pypdf = pytest.importorskip("pypdf")
    chemin = tmp_path / "essai.pdf"
    export.exporter_pdf(RECHARGES, chemin, "Tout - 2025",
                        complements=_avec_essence("Tout", "2025"))
    texte = " ".join(p.extract_text() for p in pypdf.PdfReader(chemin).pages)
    assert "voiture essence" in texte and "70,00 EUR" in texte
    assert "40 %" in texte


def test_excel_traite_les_champs_textuels_comme_du_texte(tmp_path):
    recharge = {
        **RECHARGES[1],
        "operateur": "=1+1",
        "carte": "+SUM(1,1)",
        "lieu_detail": "@SUM(1,1)",
    }
    chemin = tmp_path / "texte.xlsx"
    export.exporter_excel([recharge], chemin, "Tout")
    feuille = load_workbook(chemin, data_only=False)["Recharges"]

    for colonne, valeur in ((6, "=1+1"), (7, "+SUM(1,1)"),
                            (8, "@SUM(1,1)")):
        cellule = feuille.cell(row=2, column=colonne)
        assert cellule.value == valeur
        assert cellule.data_type == "s"
