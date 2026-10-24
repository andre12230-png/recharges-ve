"""Export des recharges en Excel (.xlsx) et PDF."""

from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape


def _fmt_date(iso: str) -> str:
    try:
        return datetime.fromisoformat(iso).strftime("%d/%m/%Y %H:%M")
    except Exception:
        return iso


def _date_excel(iso: str):
    """Vraie date pour Excel (et non du texte), afin que le tri et les
    filtres fonctionnent. Retombe sur le texte si la date est illisible."""
    try:
        return datetime.fromisoformat(iso)
    except Exception:
        return iso


def _fr(x: float, decimales: int = 2) -> str:
    """Un nombre ecrit a la francaise pour le PDF : « 1 234,50 »."""
    return f"{x:,.{decimales}f}".replace(",", "\xa0").replace(".", ",")


def _jour(iso: str) -> str:
    """« 2025-03-01T08:00 » -> « 01/03/2025 »."""
    return _fmt_date(iso)[:10]


def lignes_detail(recharges: list[dict]) -> list[list[str]]:
    """Le tableau « Détail des recharges » du PDF, entete compris, de la
    plus ancienne a la plus recente."""
    lignes = [["Date", "Lieu", "kWh", "Coût (EUR)", "EUR/kWh",
               "Tarif / Opérateur"]]
    for r in sorted(recharges, key=lambda x: x.get("date", "")):
        maison = r.get("lieu") == "maison"
        kwh = r.get("kwh", 0)
        prix = r.get("prix_kwh") or (r.get("cout_eur", 0) / kwh if kwh else 0)
        lignes.append([
            _fmt_date(r.get("date", "")),
            "Maison" if maison else "Extérieur",
            _fr(kwh),
            _fr(r.get("cout_eur", 0)),
            _fr(prix, 4),
            (r.get("tarif_nom") if maison else r.get("operateur")) or "",
        ])
    return lignes


def _stats(recharges: list[dict]) -> dict:
    """Calcule les agrégats utilisés par les deux exports."""
    maison = [r for r in recharges if r.get("lieu") == "maison"]
    ext = [r for r in recharges if r.get("lieu") == "exterieur"]

    def somme(rs, cle):
        return sum(r.get(cle, 0) for r in rs)

    par_tarif: dict[str, dict] = defaultdict(lambda: {"n": 0, "kwh": 0.0, "cout": 0.0})
    for r in maison:
        a = par_tarif[r.get("tarif_nom", "?")]
        a["n"] += 1
        a["kwh"] += r.get("kwh", 0)
        a["cout"] += r.get("cout_eur", 0)

    par_op: dict[str, dict] = defaultdict(lambda: {"n": 0, "kwh": 0.0, "cout": 0.0})
    for r in ext:
        a = par_op[r.get("operateur") or "Inconnu"]
        a["n"] += 1
        a["kwh"] += r.get("kwh", 0)
        a["cout"] += r.get("cout_eur", 0)

    return {
        "n": len(recharges),
        "kwh": somme(recharges, "kwh"),
        "cout": somme(recharges, "cout_eur"),
        "maison": {"n": len(maison), "kwh": somme(maison, "kwh"),
                   "cout": somme(maison, "cout_eur")},
        "ext": {"n": len(ext), "kwh": somme(ext, "kwh"),
                "cout": somme(ext, "cout_eur")},
        "par_tarif": dict(par_tarif),
        "par_op": dict(par_op),
    }


# ---------- Export Excel ----------

def exporter_excel(recharges: list[dict], chemin: str | Path,
                   periode_label: str = "Toute la période",
                   complements: dict | None = None) -> Path:
    """`complements` : l'abonnement et les km de la periode, calcules par la
    page Exporter (voir pages/exporter.complements) ; absent, la synthese
    s'en passe."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    chemin = Path(chemin)
    wb = Workbook()

    # --- Feuille 1 : Recharges ---
    ws = wb.active
    ws.title = "Recharges"

    entetes = ["Date", "Lieu", "kWh", "Coût (EUR)", "EUR/kWh",
               "Tarif / Opérateur", "Carte", "Détail lieu", "kWh HC", "kWh HP"]
    ws.append(entetes)

    # Bleu fonce du PDF : le blanc y fait 5,24 pour 1 (4,22 sur le bleu
    # clair d'avant, relecture du 14/10/2026).
    bleu = PatternFill("solid", fgColor="2A6BC5")
    blanc_gras = Font(bold=True, color="FFFFFF")
    for col, _ in enumerate(entetes, 1):
        c = ws.cell(row=1, column=col)
        c.fill = bleu
        c.font = blanc_gras
        c.alignment = Alignment(horizontal="center")

    # Le numero de ligne est tenu ici : demander a la feuille sa derniere
    # ligne (ws.max_row) la fait relire en entier a chaque fois, et l'export
    # de 10 000 recharges prenait 43 s (second audit du 16/10/2026).
    ligne_ecrite = 1                     # la ligne des titres
    for r in sorted(recharges, key=lambda x: x.get("date", "")):
        ligne_ecrite += 1
        tarif_op = (r.get("tarif_nom") if r.get("lieu") == "maison"
                    else r.get("operateur", ""))
        ws.append([
            _date_excel(r.get("date", "")),
            "Maison" if r.get("lieu") == "maison" else "Extérieur",
            round(r.get("kwh", 0), 3),
            round(r.get("cout_eur", 0), 2),
            round(r.get("prix_kwh", 0), 4),
            tarif_op or "",
            r.get("carte", ""),
            r.get("lieu_detail", ""),
            round(r.get("kwh_hc", 0), 2) if r.get("lieu") == "maison" else "",
            round(r.get("kwh_hp", 0), 2) if r.get("lieu") == "maison" else "",
        ])
        # Ces colonnes contiennent des textes fournis par l'utilisateur.
        # Une valeur qui commence par « = » ne doit jamais devenir une formule.
        for colonne in (6, 7, 8):
            cellule = ws.cell(row=ligne_ecrite, column=colonne)
            if isinstance(cellule.value, str):
                cellule.data_type = "s"

    # formats d'affichage : la date doit se lire comme une date, les montants
    # avec deux decimales, le prix unitaire avec quatre.
    derniere = ligne_ecrite
    for ligne in range(2, derniere + 1):
        ws.cell(row=ligne, column=1).number_format = "DD/MM/YYYY HH:MM"
        ws.cell(row=ligne, column=3).number_format = "0.000"
        ws.cell(row=ligne, column=4).number_format = "0.00"
        ws.cell(row=ligne, column=5).number_format = "0.0000"
        ws.cell(row=ligne, column=9).number_format = "0.00"
        ws.cell(row=ligne, column=10).number_format = "0.00"

    # ligne de totaux, en gras sous le tableau
    if derniere >= 2:
        total = derniere + 1
        ws.cell(row=total, column=1, value="TOTAL").font = Font(bold=True)
        for col in (3, 4, 9, 10):
            lettre = get_column_letter(col)
            c = ws.cell(row=total, column=col,
                        value=f"=SUM({lettre}2:{lettre}{derniere})")
            c.font = Font(bold=True)
            c.number_format = "0.000" if col == 3 else "0.00"

    # largeurs de colonnes
    largeurs = [17, 11, 9, 11, 10, 24, 13, 36, 9, 9]
    for i, w in enumerate(largeurs, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    # filtres sur l'entete : tri par date, par lieu, par operateur...
    ws.auto_filter.ref = f"A1:{get_column_letter(len(entetes))}{derniere}"

    # --- Feuille 2 : Synthese ---
    ws2 = wb.create_sheet("Synthèse")
    st = _stats(recharges)
    titre_font = Font(bold=True, size=13, color="2A6BC5")   # 5,24 sur le blanc

    ws2["A1"] = "SYNTHÈSE DES RECHARGES"
    ws2["A1"].font = titre_font
    ws2["A2"] = f"Période : {periode_label}"
    ws2["A3"] = f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}"

    lignes = [
        ("", ""),
        ("GLOBAL", ""),
        ("Nombre de recharges", st["n"]),
        ("Énergie totale (kWh)", round(st["kwh"], 2)),
        ("Coût total (EUR)", round(st["cout"], 2)),
        ("Prix moyen (EUR/kWh)",
         round(st["cout"] / st["kwh"], 4) if st["kwh"] else 0),
        ("", ""),
        ("MAISON", ""),
        ("Recharges", st["maison"]["n"]),
        ("kWh", round(st["maison"]["kwh"], 2)),
        ("Coût (EUR)", round(st["maison"]["cout"], 2)),
        ("", ""),
        ("EXTÉRIEUR", ""),
        ("Recharges", st["ext"]["n"]),
        ("kWh", round(st["ext"]["kwh"], 2)),
        ("Coût (EUR)", round(st["ext"]["cout"], 2)),
    ]
    complements = complements or {}
    if complements.get("abonnement") is not None:
        lignes += [
            ("", ""),
            ("ABONNEMENT", ""),
            ("Abonnement électricité de la période (EUR)",
             round(complements["abonnement"], 2)),
            ("À part : il n'entre ni dans le coût des recharges ni dans le "
             "prix du kWh.", ""),
        ]
    if complements.get("frais_cartes") is not None:
        lignes += [
            ("", ""),
            ("CARTES DE RECHARGE", ""),
            ("Frais de cartes de recharge de la période (EUR)",
             round(complements["frais_cartes"], 2)),
            ("Achats de cartes et abonnements : à part, eux aussi.", ""),
        ]
    km = complements.get("km")
    if km:
        lignes += [
            ("", ""),
            ("KILOMÉTRAGE", ""),
            ("Entre vos relevés du compteur",
             f"du {_jour(km['debut'])} au {_jour(km['fin'])}"),
            ("Distance parcourue (km)", km["km"]),
            ("Coût des recharges (EUR)", km["cout"]),
            ("Coût aux 100\xa0km (EUR)", km["cout_100"]),
        ]
    essence = complements.get("essence")
    if essence:
        lignes += [
            ("", ""),
            ("COMPARAISON AVEC UNE VOITURE ESSENCE", ""),
            ("Kilomètres comparés",
             f"{essence['km']:,.0f}\xa0km, du {_jour(essence['debut'])} au "
             f"{_jour(essence['fin'])}".replace(",", "\xa0")),
            ("Voiture essence (L/100\xa0km)", essence["conso_essence"]),
            ("Prix moyen du litre (EUR)", essence["prix_essence"]),
            ("En essence (EUR)", essence["cout_essence"]),
            ("En électricité (EUR)", essence["cout_electrique"]),
            ("Économie (EUR)", essence["economie"]),
        ]
        if essence.get("part_non_enregistree"):
            lignes.append((texte_part_non_enregistree(essence), ""))
    for label, val in lignes:
        ws2.append([label, val])
        # gras sur les titres de section, au moment ou la ligne est ecrite
        if label in ("GLOBAL", "MAISON", "EXTÉRIEUR", "ABONNEMENT",
                     "CARTES DE RECHARGE", "KILOMÉTRAGE",
                     "COMPARAISON AVEC UNE VOITURE ESSENCE"):
            ws2.cell(row=ws2.max_row, column=1).font = Font(bold=True)

    ws2.column_dimensions["A"].width = 42
    for lettre in ("B", "C", "D", "E"):
        ws2.column_dimensions[lettre].width = 16

    def _tableau(titre: str, donnees: dict):
        """Ecrit un bloc « repartition » avec son entete en gras."""
        if not donnees:
            return
        ws2.append(["", ""])
        ws2.append([titre, "Recharges", "kWh", "Coût (EUR)", "EUR/kWh"])
        for c in ws2[ws2.max_row]:
            c.font = Font(bold=True)
        for nom, a in sorted(donnees.items(), key=lambda x: -x[1]["cout"]):
            moy = a["cout"] / a["kwh"] if a["kwh"] else 0
            ws2.append([nom, a["n"], round(a["kwh"], 2), round(a["cout"], 2),
                        round(moy, 4)])
            ligne = ws2.max_row
            ws2.cell(row=ligne, column=3).number_format = "0.00"
            ws2.cell(row=ligne, column=4).number_format = "0.00"
            ws2.cell(row=ligne, column=5).number_format = "0.0000"

    # repartition par tarif (maison) : elle manquait dans le fichier Excel
    # alors qu'elle figure dans le PDF et dans la vue Statistiques
    _tableau("PAR TARIF (maison)", st["par_tarif"])
    _tableau("PAR OPÉRATEUR (extérieur)", st["par_op"])

    # La Synthèse ne contient aucune formule : tout texte y reste un texte.
    # Un nom de tarif ou d'opérateur « =1+1 » devenait sinon une formule
    # Excel (audit du 14/10/2026 ; la feuille Recharges était déjà protégée).
    for ligne in ws2.iter_rows():
        for cellule in ligne:
            if isinstance(cellule.value, str):
                cellule.data_type = "s"

    wb.save(chemin)
    return chemin


# ---------- Export PDF ----------

def exporter_pdf(recharges: list[dict], chemin: str | Path,
                 periode_label: str = "Toute la période",
                 complements: dict | None = None) -> Path:
    """Rapport imprimable : synthese, abonnement et km s'ils sont fournis
    (voir exporter_excel), repartitions, puis la liste des recharges."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer)
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    chemin = Path(chemin)
    st = _stats(recharges)
    styles = getSampleStyleSheet()
    # Titre au bleu fonce des sous-titres : 5,24 pour 1 sur le blanc ; le
    # bleu clair d'avant (#3A7BD5) ne faisait que 4,22 (relecture du 14/10/2026).
    titre = ParagraphStyle("titre", parent=styles["Title"],
                           textColor=colors.HexColor("#2A6BC5"))
    h2 = ParagraphStyle("h2", parent=styles["Heading2"],
                        textColor=colors.HexColor("#2A6BC5"))

    doc = SimpleDocTemplate(str(chemin), pagesize=A4,
                            topMargin=1.5 * cm, bottomMargin=1.5 * cm,
                            leftMargin=1.5 * cm, rightMargin=1.5 * cm)
    elements = []

    elements.append(Paragraph("Rapport des recharges VE", titre))
    elements.append(Paragraph(
        f"Période : {_texte_pdf(periode_label)}<br/>"
        f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}",
        styles["Normal"]))
    elements.append(Spacer(1, 0.6 * cm))

    # --- Synthese globale ---
    prix_moyen = st["cout"] / st["kwh"] if st["kwh"] else 0
    data_syn = [
        ["", "Recharges", "Énergie (kWh)", "Coût (EUR)", "EUR/kWh"],
        ["Total", st["n"], _fr(st["kwh"], 1), _fr(st["cout"]),
         _fr(prix_moyen, 4)],
        ["Maison", st["maison"]["n"], _fr(st["maison"]["kwh"], 1),
         _fr(st["maison"]["cout"]),
         _fr(st["maison"]["cout"] / st["maison"]["kwh"], 4)
         if st["maison"]["kwh"] else "-"],
        ["Extérieur", st["ext"]["n"], _fr(st["ext"]["kwh"], 1),
         _fr(st["ext"]["cout"]),
         _fr(st["ext"]["cout"] / st["ext"]["kwh"], 4)
         if st["ext"]["kwh"] else "-"],
    ]
    t = Table(data_syn, colWidths=[3 * cm, 3 * cm, 3.5 * cm, 3 * cm, 3 * cm])
    t.setStyle(TableStyle([
        # Bleu fonce des autres en-tetes : le blanc fait 5,24 pour 1 ; sur
        # le bleu clair d'avant (#3A7BD5), 4,22 (relecture du 14/10/2026).
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2A6BC5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (0, -1), "Helvetica-Bold"),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.white, colors.HexColor("#EEF3FB")]),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    elements.append(t)
    elements.append(Spacer(1, 0.6 * cm))

    # --- Abonnement et kilometrage (28/09/2026) ---
    complements = complements or {}
    if complements.get("abonnement") is not None:
        elements.append(Paragraph(
            f"<b>Abonnement électricité de la période :</b> "
            f"{_fr(complements['abonnement'])}\xa0EUR. À part : il n'entre ni "
            "dans le coût des recharges ni dans le prix du kWh.",
            styles["Normal"]))
        elements.append(Spacer(1, 0.3 * cm))
    if complements.get("frais_cartes") is not None:
        elements.append(Paragraph(
            f"<b>Frais de cartes de recharge de la période :</b> "
            f"{_fr(complements['frais_cartes'])}\xa0EUR (achats de cartes et "
            "abonnements), à part eux aussi.", styles["Normal"]))
        elements.append(Spacer(1, 0.3 * cm))
    km = complements.get("km")
    if km:
        elements.append(Paragraph(
            f"<b>Coût de vos recharges aux 100\xa0km :</b> "
            f"{_fr(km['cout_100'])}\xa0EUR/100\xa0km — {_fr(km['cout'])}\xa0EUR de "
            f"recharges enregistrées pour {_fr(km['km'], 0)}\xa0km, du "
            f"{_jour(km['debut'])} au {_jour(km['fin'])}, entre vos relevés "
            "du compteur (abonnement et frais de cartes non compris).",
            styles["Normal"]))
    essence = complements.get("essence")
    if essence:
        texte = (f"<b>Comparaison avec une voiture essence :</b> "
                 f"{_fr(essence['km'], 0)}\xa0km, du {_jour(essence['debut'])} au "
                 f"{_jour(essence['fin'])}, à "
                 f"{_fr(essence['conso_essence'], 1)}\xa0L/100\xa0km et "
                 f"{_fr(essence['prix_essence'])}\xa0EUR/L en moyenne : "
                 f"{_fr(essence['cout_essence'])}\xa0EUR en essence, contre "
                 f"{_fr(essence['cout_electrique'])}\xa0EUR en électricité, soit "
                 f"{_fr(essence['economie'])}\xa0EUR d'économie.")
        if essence.get("part_non_enregistree"):
            texte += " " + texte_part_non_enregistree(essence)
        elements.append(Spacer(1, 0.3 * cm))
        elements.append(Paragraph(texte, styles["Normal"]))
    if (complements.get("abonnement") is not None or km or essence
            or complements.get("frais_cartes") is not None):
        elements.append(Spacer(1, 0.6 * cm))

    # --- Par tarif (maison) ---
    if st["par_tarif"]:
        elements.append(Paragraph("Répartition par tarif (maison)", h2))
        data_t = [["Tarif", "Recharges", "kWh", "Coût (EUR)", "EUR/kWh"]]
        for nom, a in st["par_tarif"].items():
            moy = a["cout"] / a["kwh"] if a["kwh"] else 0
            data_t.append([nom, a["n"], _fr(a["kwh"], 1),
                           _fr(a["cout"]), _fr(moy, 4)])
        elements.append(_table_simple(data_t))
        elements.append(Spacer(1, 0.6 * cm))

    # --- Par operateur (exterieur) ---
    if st["par_op"]:
        elements.append(Paragraph("Répartition par opérateur (extérieur)", h2))
        data_o = [["Opérateur", "Recharges", "kWh", "Coût (EUR)", "EUR/kWh"]]
        for nom, a in sorted(st["par_op"].items(), key=lambda x: -x[1]["cout"]):
            moy = a["cout"] / a["kwh"] if a["kwh"] else 0
            data_o.append([nom, a["n"], _fr(a["kwh"], 1),
                           _fr(a["cout"]), _fr(moy, 4)])
        elements.append(_table_simple(data_o))
        elements.append(Spacer(1, 0.6 * cm))

    # --- Detail des recharges : sans lui, le PDF n'etait qu'une synthese ---
    elements.append(Paragraph("Détail des recharges", h2))
    petit = ParagraphStyle("petit", parent=styles["Normal"], fontSize=8,
                           leading=10)
    detail = lignes_detail(recharges)
    # Le nom du tarif ou de l'operateur peut etre long : il passe a la ligne
    # dans sa case au lieu de deborder.
    for ligne in detail[1:]:
        ligne[5] = Paragraph(_texte_pdf(ligne[5]), petit)
    t = Table(detail, repeatRows=1,       # entete repete sur chaque page
              colWidths=[3.2 * cm, 2.2 * cm, 1.9 * cm, 2.2 * cm, 2.1 * cm,
                         6.4 * cm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2A6BC5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (2, 0), (4, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.white, colors.HexColor("#EEF3FB")]),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
    ]))
    elements.append(t)

    doc.build(elements)
    return chemin


def texte_part_non_enregistree(essence: dict) -> str:
    """La phrase de l'export sur l'énergie absente des recharges.

    C'est une estimation, tirée de la consommation affichée par la voiture :
    l'origine et le prix de cette énergie ne sont pas connus. Le texte disait
    « elle n'a rien coûté », ce que les relevés ne permettent pas d'affirmer
    (audit du 14/10/2026)."""
    return (f"Environ {essence['part_non_enregistree']}\xa0% de l'énergie de "
            "ces kilomètres ne figure pas dans les recharges (estimation "
            "d'après la consommation affichée par la voiture) : recharges "
            "chez des proches, bornes gratuites ou recharges non saisies. "
            "Son origine et son prix ne sont pas connus ; le coût en "
            "électricité ne compte que les recharges saisies.")


def _texte_pdf(texte: str) -> str:
    """Un texte saisi par l'utilisateur, prêt pour un paragraphe du PDF.

    ReportLab lit les paragraphes comme du balisage : un nom d'opérateur
    « Borne <b> » faisait échouer l'export, et « <i>Rapide</i> » passait en
    italique (audit du 14/10/2026). Les signes < > & sont donc échappés : le
    nom s'affiche tel quel. Les balises écrites par l'application elle-même
    ne passent pas par ici."""
    return escape(str(texte))


def _table_simple(data):
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import Table, TableStyle

    t = Table(data, colWidths=[5.5 * cm, 2.5 * cm, 2.5 * cm, 3 * cm, 2.5 * cm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2A6BC5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.white, colors.HexColor("#EEF3FB")]),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
    ]))
    return t
